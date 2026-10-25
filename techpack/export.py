"""Editable exports. Factories dislike locked PDFs, so packs also come as Excel."""

import io
import re

from .schema import TechPack


def grade_steps(values) -> str:
    """Grade increments between consecutive sizes, e.g. '+3 per size' or '+2 / +3'."""
    nums = []
    for v in values:
        m = re.search(r"-?\d+(?:\.\d+)?", str(v))
        if not m:
            return "–"
        nums.append(float(m.group()))
    if len(nums) < 2:
        return "–"
    steps = [round(b - a, 2) for a, b in zip(nums, nums[1:])]
    if len(set(steps)) == 1:
        return f"{steps[0]:+g} per size"
    return " / ".join(f"{d:+g}" for d in steps)


def build_xlsx(pack: TechPack, meta: dict, version: int) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    brief = meta.get("brief", {})
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="EEEEEE")

    def sheet(title, header, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = head_fill
        for r in rows:
            ws.append([str(c) for c in r])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(12, max(len(str(c.value or "")) for c in col) + 2))
        ws.freeze_panes = "A2"

    sheet("Cover", ["Field", "Value"], [
        ("Style name", pack.style_name), ("Style number", pack.style_number), ("Version", f"v{version}"),
        ("Brand", brief.get("brand", "")), ("Factory / supplier", brief.get("factory_or_supplier", "")),
        ("Season", pack.season), ("Category", pack.category), ("Fit", pack.gender_fit),
        ("Sample stage", brief.get("sample_stage", "")), ("Order quantity", brief.get("order_quantity", "")),
        ("Target delivery", brief.get("target_delivery", "")), ("Target cost", brief.get("target_cost", "")),
        ("Units", pack.units), ("Sizes", ", ".join(pack.sizes)), ("Base size", pack.base_size),
        ("Country of origin", pack.country_of_origin), ("Summary", pack.summary),
    ], first=True)
    sheet("BOM", ["Component", "Material", "Composition", "Weight/size", "Colour", "Pantone", "Placement", "Qty", "Notes"],
          [(b.component, b.material, b.composition, b.weight_or_size, b.color, b.pantone, b.placement, b.quantity, b.notes)
           for b in pack.bill_of_materials])
    sheet("Colourways", ["Colourway", "Pantone", "Applies to"], [(c.name, c.pantone, c.applies_to) for c in pack.colorways])
    sheet("Measurements", ["Code", "Point of measure", "Tol +/-", "Grade"] + pack.sizes,
          [(m.code, m.description, m.tolerance, grade_steps(m.values), *m.values) for m in pack.measurements])
    sheet("Construction", ["Area", "Operation", "Stitch", "Seam", "SPI", "Seam allowance", "Callout", "Notes"],
          [(c.area, c.operation, c.stitch_type, c.seam_type, c.spi, c.seam_allowance, c.callout_ref, c.notes)
           for c in pack.construction])
    sheet("Artwork", ["Artwork", "Technique", "Size", "Placement", "Colours", "Notes"],
          [(a.name, a.technique, a.size, a.placement, a.colors, a.notes) for a in pack.artwork])
    sheet("Labels & trims", ["Item", "Spec", "Size", "Content", "Placement"],
          [(l.item, l.spec, l.size, l.content, l.placement) for l in pack.labels_and_trims])
    sheet("Testing", ["Test", "Standard", "Requirement"],
          [(t.test, t.standard, t.requirement) for t in pack.testing_requirements])
    pk = pack.packing_details
    sheet("Packing", ["Item", "Specification"], [
        ("Folding", pk.folding), ("Polybag", pk.polybag), ("Hangtag / stickers", pk.hangtag_and_stickers),
        ("Units per carton", pk.units_per_carton), ("Packing ratio", pk.packing_ratio), ("Carton marking", pk.carton_marking),
    ])
    sheet("Callouts", ["#", "Detail", "Description"], [(c.number, c.label, c.description) for c in pack.callouts])
    sheet("Revision log", ["Version", "Date", "Note"], [(h["version"], h["at"], h["note"]) for h in meta.get("history", [])])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
